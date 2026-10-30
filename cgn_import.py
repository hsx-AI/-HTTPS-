"""Import CGN exports without discarding source-specific or repeated columns."""
from __future__ import annotations
import hashlib
import json
import sqlite3
from collections import Counter
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from excel_db import EXCEL_COLUMNS, ensure_schema

MAPPING = {'文件名称': '质量计划名称', '文件版本': '版本', '所属项目': '项目', '制造单位': '供应商',
           '流程状态': '状态', '创建时间': '创建日期', '发起人': '创建人', '适用机组号': '机组',
           '供应商名称': '供应商', '设备名称': '适用物项名称', '合同号': '合同订单编号', '审查结论': '审查结论'}

def text(value):
    if value is None: return ''
    if isinstance(value, datetime): return value.isoformat(sep=' ')
    if isinstance(value, date): return value.isoformat()
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()

def read_rows(path):
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            # The platform export advertises only two rows despite containing hundreds.
            # Stream actual worksheet cells instead of trusting its dimension metadata.
            sheet.reset_dimensions()
            rows = [[text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        finally: book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(path)
        try:
            sheet = book.sheet_by_index(0)
            rows = [[text(xlrd.xldate_as_datetime(c.value, book.datemode)) if c.ctype == xlrd.XL_CELL_DATE else text(c.value) for c in sheet.row(i)] for i in range(sheet.nrows)]
        finally: book.release_resources()
    if len(rows) < 2: raise ValueError('广核导出表格为空，本地已有数据已保留')
    headers = rows[0]
    for field in ('质量计划名称', '质量计划编码', '版本'):
        if headers.count(field) != 1: raise ValueError(f'广核导出缺少或重复关键列：{field}')
    records, seen = [], set()
    for line, values in enumerate(rows[1:], 2):
        if not any(values): continue
        values += [''] * (len(headers)-len(values))
        first = {h: values[i] for i, h in reversed(list(enumerate(headers)))}
        code, version = first['质量计划编码'], first['版本']
        if not code: raise ValueError(f'第 {line} 行没有质量计划编码，已停止导入')
        key = 'cgn:' + hashlib.sha256(json.dumps([code, version], ensure_ascii=False).encode()).hexdigest()
        if key in seen: raise ValueError(f'第 {line} 行的计划编码和版本重复，需核对后导入')
        seen.add(key)
        counts = Counter(); fields = []
        for i, header in enumerate(headers):
            counts[header] += 1
            label = ('流程状态' if counts[header] == 1 else '版本状态' if counts[header] == 2 else f'状态（第{counts[header]}处）') if header == '状态' else header or f'第{i+1}列'
            fields.append({'name': header, 'label': label, 'column': i+1, 'value': values[i]})
        record = {k: first.get(v, '') for k,v in MAPPING.items()}
        record['文件编号'] = key
        record['版本状态'] = next((f['value'] for f in fields if f['label'] == '版本状态'), '')
        state = record['流程状态']
        record.update(display_code=code, raw_data=json.dumps(fields, ensure_ascii=False),
                      completion_state='done' if state in ('完成','已完成') else 'active' if state in ('审查中','编制中','审核中','审批中') else 'unknown',
                      内部编号=first.get('内部编码',''))
        records.append(record)
    if not records: raise ValueError('没有可导入的计划，本地已有数据已保留')
    return records

def import_excel(db_path: Path, excel_path: Path, expected_count=None):
    records = read_rows(excel_path)  # Validate every row before touching the database.
    if expected_count is not None and len(records) != expected_count: raise ValueError(f'导出数量与查询结果不同（{len(records)} / {expected_count}），已保留旧数据')
    db_path.parent.mkdir(parents=True, exist_ok=True)
    now = datetime.now().astimezone().isoformat(timespec='seconds'); batch = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    inserted = updated = changed = 0
    extras = ['display_code','raw_data','completion_state','内部编号']
    with closing(sqlite3.connect(db_path, timeout=30)) as conn:
        conn.row_factory = sqlite3.Row
        ensure_schema(conn)
        columns = {r[1] for r in conn.execute('PRAGMA table_info(documents)')}
        for k in extras:
            if k not in columns: conn.execute(f'ALTER TABLE documents ADD COLUMN "{k}" TEXT')
        conn.execute("CREATE TABLE IF NOT EXISTS source_info (source_id TEXT PRIMARY KEY)")
        conn.execute("INSERT OR IGNORE INTO source_info VALUES ('cgn')")
        conn.commit()
        with conn:
            for record in records:
                key = record['文件编号']; old = conn.execute('SELECT * FROM documents WHERE "文件编号"=?', [key]).fetchone()
                if old is not None and old['raw_data'] != record['raw_data']:
                    before = {(x['name'],x['column']):x for x in json.loads(old['raw_data'])}
                    after = {(x['name'],x['column']):x for x in json.loads(record['raw_data'])}
                    for field in before.keys() | after.keys():
                        a,b = before.get(field,{}),after.get(field,{})
                        if a.get('value','') != b.get('value',''):
                            conn.execute('INSERT INTO document_changes ("文件编号",field_name,old_value,new_value,changed_at,source_file,sync_batch_id) VALUES (?,?,?,?,?,?,?)', [key,b.get('label') or a['label'],a.get('value',''),b.get('value',''),now,excel_path.name,batch]); changed += 1
                    updated += 1
                if old is None: inserted += 1
                cols = EXCEL_COLUMNS + extras + ['fetched_at','updated_at','source_file']
                data = [record.get(k,'') for k in EXCEL_COLUMNS+extras]+[now, now if old is None or old['raw_data']!=record['raw_data'] else old['updated_at'], excel_path.name]
                conn.execute('INSERT OR REPLACE INTO documents ('+','.join('"'+k+'"' for k in cols)+') VALUES ('+','.join('?' for _ in cols)+')', data)
    return {'total':len(records),'inserted':inserted,'updated':updated,'changed_fields':changed}
