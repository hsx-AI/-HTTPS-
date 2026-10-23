"""User-maintained plan baseline, stored separately from received AE snapshots."""
from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
from collections import Counter, defaultdict
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from .db import resolve_db_path

MAX_ROWS = 10000
FIELDS = {
    'name': ('计划名称', ['计划名称', '文件名称', '质量计划名称']),
    'code': ('文件编号', ['文件编号', '外部编号', '质量计划编号']),
    'internalCode': ('内部编号', ['内部编号', '内部文件编号', '内部编码']),
    'project': ('所属项目', ['所属项目', '项目', '项目名称']),
    'unit': ('适用机组号', ['适用机组号', '机组', '机组号']),
    'company': ('制造单位', ['制造单位', '供方', '供应商名称']),
    'plannedDate': ('计划完成时间', ['计划完成时间', '计划完成日期', '计划日期']),
    'latestApprovalAt': ('最近审批时间', ['最近审批流程时间', '最新审批流程时间', '最新审批时间', '最近审批时间']),
}


def catalog_path() -> Path:
    return Path(os.environ.get('PLAN_CATALOG_PATH') or resolve_db_path().parent / 'quality_plan_catalog.db')


def text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value) -> str:
    return ''.join(text(value).split()).casefold()


def ensure_schema(conn):
    conn.execute('CREATE TABLE IF NOT EXISTS plans (id INTEGER PRIMARY KEY, identity TEXT UNIQUE NOT NULL, ' + ', '.join(f'"{k}" TEXT NOT NULL DEFAULT \'\'' for k in [*FIELDS, 'filename', 'importedAt']) + ')')
    conn.execute('CREATE TABLE IF NOT EXISTS document_notes (code TEXT PRIMARY KEY, internal_code TEXT NOT NULL, updated_at TEXT NOT NULL)')


def save_note(code: str, internal_code: str) -> dict:
    if not isinstance(code, str) or not isinstance(internal_code, str) or not code.strip() or len(internal_code) > 200:
        raise ValueError('备注格式无效，内部编号最多 200 字符')
    path = catalog_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    now = datetime.now().astimezone().isoformat(timespec='seconds')
    with closing(sqlite3.connect(path)) as conn, conn:
        ensure_schema(conn)
        conn.execute('INSERT INTO document_notes VALUES (?, ?, ?) ON CONFLICT(code) DO UPDATE SET internal_code=excluded.internal_code, updated_at=excluded.updated_at', [code, internal_code.strip(), now])
    return {'code': code, 'internalCode': internal_code.strip(), 'updatedAt': now}


def read_notes() -> dict:
    path = catalog_path()
    if not path.exists():
        return {}
    with closing(sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)) as conn:
        if not conn.execute("SELECT 1 FROM sqlite_master WHERE name='document_notes'").fetchone():
            return {}
        return dict(conn.execute('SELECT code, internal_code FROM document_notes'))


def read_catalog() -> list[dict]:
    path = catalog_path()
    if not path.exists():
        return []
    with closing(sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)) as conn:
        conn.row_factory = sqlite3.Row
        return [dict(row) for row in conn.execute('SELECT * FROM plans ORDER BY id')]


def preview_file(payload: bytes, filename: str, sheet: str = '') -> dict:
    suffix = Path(filename).suffix.lower()
    sheets = []
    if suffix == '.csv':
        try:
            decoded = payload.decode('utf-8-sig')
        except UnicodeDecodeError:
            decoded = payload.decode('gb18030')
        reader = csv.reader(io.StringIO(decoded))
        rows = []
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_ROWS + 31:
                raise ValueError('清单最多支持 10000 条记录')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        try:
            with ZipFile(io.BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                    raise ValueError('Excel 解压后的内容过大')
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        except (BadZipFile, KeyError) as exc:
            raise ValueError('无法读取 Excel，请使用有效的 .xlsx 文件') from exc
        try:
            sheets = book.sheetnames
            if sheet and sheet not in sheets:
                raise ValueError('所选工作表不存在')
            sheet = sheet or sheets[0]
            ws = book[sheet]
            if ws.max_column and ws.max_column > 100:
                raise ValueError('清单最多支持 100 列')
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([text(v) for v in row])
                if len(rows) > MAX_ROWS + 31:
                    raise ValueError('清单最多支持 10000 条记录')
        finally:
            book.close()
    elif suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=payload)
        try:
            sheets = book.sheet_names()
            if sheet and sheet not in sheets:
                raise ValueError('所选工作表不存在')
            sheet = sheet or sheets[0]
            ws = book.sheet_by_name(sheet)
            if ws.nrows > MAX_ROWS + 31 or ws.ncols > 100:
                raise ValueError('清单最多支持 10000 条记录、100 列')
            rows = [[text(xlrd.xldate_as_datetime(c.value, book.datemode)) if c.ctype == xlrd.XL_CELL_DATE else text(c.value) for c in ws.row(i)] for i in range(ws.nrows)]
        finally:
            book.release_resources()
    else:
        raise ValueError('请上传 .xlsx、.xls 或 .csv 文件')
    aliases = {norm(alias) for _, names in FIELDS.values() for alias in names}
    header_index = next((i for i, row in enumerate(rows[:30]) if sum(norm(v) in aliases for v in row) >= 2), 0)
    if not rows:
        raise ValueError('文件为空')
    headers = [text(v) or f'未命名列{i + 1}' for i, v in enumerate(rows[header_index])]
    if len(headers) > 100:
        raise ValueError('清单最多支持 100 列')
    data = [[text(v) for v in row[:len(headers)]] for row in rows[header_index + 1:] if any(text(v) for v in row)]
    if not data or len(data) > MAX_ROWS:
        raise ValueError('清单应包含 1 至 10000 条记录')
    mapping = {key: next((i for i, h in enumerate(headers) if norm(h) in {norm(a) for a in aliases}), -1) for key, (_, aliases) in FIELDS.items()}
    return {'headers': headers, 'rows': data, 'mapping': mapping, 'fields': [{'key': k, 'label': v[0]} for k, v in FIELDS.items()], 'sheets': sheets, 'sheet': sheet, 'headerRow': header_index + 1}


def import_rows(rows: list, mapping: dict, filename: str) -> dict:
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
        raise ValueError('清单应包含 1 至 10000 条记录')
    if not isinstance(mapping, dict) or any(type(i) is not int or i < -1 or i > 99 for i in mapping.values()):
        raise ValueError('字段映射无效')
    used = [mapping.get(k, -1) for k in FIELDS if mapping.get(k, -1) >= 0]
    if len(set(used)) != len(used):
        raise ValueError('同一列表格内容不能重复用于不同信息，请检查上面的选择')
    if mapping.get('name', -1) < 0:
        raise ValueError('请选择哪一列是计划名称')
    records, keys = [], set()
    now = datetime.now().astimezone().isoformat(timespec='seconds')
    for n, row in enumerate(rows, 1):
        if not isinstance(row, list) or len(row) > 100 or any(isinstance(v, (dict, list)) for v in row):
            raise ValueError(f'第 {n} 条记录格式无效')
        record = {k: text(row[mapping[k]]) if 0 <= mapping.get(k, -1) < len(row) else '' for k in FIELDS}
        if any(len(v) > 1000 for v in record.values()):
            raise ValueError(f'第 {n} 条记录字段超过 1000 字符')
        if not record['name'] or not (record['code'] or record['internalCode'] or (record['project'] and record['unit'])):
            raise ValueError(f'第 {n} 条记录需有计划名称，以及文件编号、内部编号或项目与机组信息')
        for k in ('plannedDate', 'latestApprovalAt'):
            if record[k]:
                try:
                    parsed = datetime.fromisoformat(record[k].replace('/', '-'))
                    record[k] = parsed.date().isoformat() if k == 'plannedDate' or len(record[k]) == 10 else parsed.isoformat(sep=' ')
                except ValueError as exc:
                    raise ValueError(f'第 {n} 条记录的{FIELDS[k][0]}格式无效，请使用 YYYY-MM-DD 或 YYYY-MM-DD HH:MM:SS') from exc
        identity = json.dumps([norm(record['project']), norm(record['unit']), 'code' if record['code'] else 'internal' if record['internalCode'] else 'name', norm(record['code'] or record['internalCode'] or record['name'])], ensure_ascii=False)
        if identity in keys:
            raise ValueError(f'第 {n} 条记录与本次清单中的其他条目重复，请先核对编号及项目机组')
        keys.add(identity)
        records.append({**record, 'identity': identity, 'filename': Path(filename).name[:255], 'importedAt': now})
    path = catalog_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(path)) as conn, conn:
        ensure_schema(conn)
        existing = {r[0] for r in conn.execute('SELECT identity FROM plans')}
        columns = list(records[0])
        names = ', '.join(f'"{k}"' for k in columns)
        updates = ', '.join(f'"{k}"=excluded."{k}"' for k in columns if k != 'identity')
        conn.executemany(f'INSERT INTO plans ({names}) VALUES ({",".join("?" for _ in columns)}) ON CONFLICT(identity) DO UPDATE SET {updates}', [[r[k] for k in columns] for r in records])
    return {'inserted': len(keys - existing), 'updated': len(keys & existing), 'total': len(records), 'importedAt': now}


def match_catalog(plans: list[dict], documents: list[dict]) -> list[dict]:
    indexes = {key: defaultdict(list) for key in ('code', 'internalCode', 'name')}
    for doc in documents:
        for key, index in indexes.items():
            # Baseline-derived values must not become evidence for another baseline match.
            if key == 'internalCode' and doc.get('internalCodeSource') == '导入清单':
                continue
            if norm(doc.get(key)):
                index[norm(doc[key])].append(doc)
    result = []
    for plan in plans:
        candidates, method = [], ''
        for key in ('code', 'internalCode', 'name'):
            value = norm(plan.get(key))
            if not value:
                continue
            # Names only identify a plan together with its project and unit.
            if key == 'name' and not (plan.get('project') and plan.get('unit')):
                continue
            candidates = [d for d in indexes[key].get(value, []) if all(not plan.get(k) or norm(plan[k]) == norm(d.get(k)) for k in ('project', 'unit'))]
            if candidates:
                method = {'code': '文件编号', 'internalCode': '内部编号', 'name': '名称＋项目＋机组'}[key]
                break
            # An explicitly supplied external number must not silently match a different number by name.
            if key == 'code':
                break
        result.append({**plan, 'matches': candidates, 'matchMethod': method})
    uses = Counter(d['code'] for row in result if len(row['matches']) == 1 for d in row['matches'])
    for row in result:
        docs = row.pop('matches')
        ambiguous = len(docs) > 1 or (len(docs) == 1 and uses[docs[0]['code']] > 1)
        row['candidateCodes'] = [d['code'] for d in docs]
        row['matchCount'] = len(docs)
        row['actual'] = docs[0] if len(docs) == 1 and not ambiguous else None
        row['comparisonStatus'] = '待核对' if ambiguous else '未接入' if not docs else '已完成' if text(docs[0]['status']) == '完成' else '处理中'
    return result


def enrich_documents(documents: list[dict], plans: list[dict] | None = None) -> list[dict]:
    notes = read_notes()
    for doc in documents:
        doc['internalCodeSource'] = '源平台' if doc.get('internalCode') else ''
        doc['latestApprovalSource'] = '源平台' if doc.get('latestApprovalAt') else ''
        if notes.get(doc['code']):
            doc['internalCode'] = notes[doc['code']]
            doc['internalCodeSource'] = '用户备注'
    for row in match_catalog(read_catalog() if plans is None else plans, documents):
        doc = row['actual']
        if doc is not None:
            for key, source_key in [('internalCode', 'internalCodeSource'), ('latestApprovalAt', 'latestApprovalSource')]:
                if not doc.get(key) and row.get(key):
                    doc[key] = row[key]
                    doc[source_key] = '导入清单'
    return documents


def compare(documents: list[dict], project='全部', group='全部', q='', status='全部', company='', page=1, page_size=20) -> dict:
    matched = match_catalog(read_catalog(), documents)
    for row in matched:
        if row['actual']:
            for key in ('project', 'unit'):
                row[key] = row[key] or row['actual'].get(key, '')
    scoped = [r for r in matched if (project == '全部' or r['project'] == project) and (group == '全部' or r['unit'] == group) and (not company or r['company'] == company or (r['actual'] and r['actual']['company'] == company))]
    summary = dict(Counter(r['comparisonStatus'] for r in scoped))
    total_plans = len(scoped)
    matched_codes = {r['actual']['code'] for r in matched if r['actual']}
    uncertain_codes = {code for r in matched if r['comparisonStatus'] == '待核对' for code in r['candidateCodes']}
    extras = [d for d in documents if d['code'] not in matched_codes | uncertain_codes and (project == '全部' or d['project'] == project) and (group == '全部' or d['unit'] == group) and (not company or d['company'] == company)]
    rows = [{**d, 'actual': d, 'comparisonStatus': '清单外记录', 'matchMethod': '', 'id': 'actual:' + d['code']} for d in extras] if status == '清单外记录' else [r for r in scoped if status == '全部' or r['comparisonStatus'] == status]
    if q.strip():
        needle = q.strip().casefold()
        rows = [r for r in rows if any(needle in text(r.get(k)).casefold() or needle in text((r.get('actual') or {}).get(k)).casefold() for k in ('name', 'code', 'internalCode', 'company'))]
    return {'items': rows[(page-1)*page_size:page*page_size], 'total': len(rows), 'page': page, 'pageSize': page_size, 'summary': summary, 'plannedTotal': total_plans, 'completionRate': round(summary.get('已完成', 0)*100/total_plans, 1) if total_plans else 0, 'extraCount': len(extras)}
