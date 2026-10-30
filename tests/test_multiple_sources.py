import json
import os
import sqlite3
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED
import re
from contextlib import closing
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient
from dashboard_api.main import app
from dashboard_api import catalog
from cgn_import import import_excel
from test_dashboard_feedback import create_snapshot

class MultiSourceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.ae = self.root/'ae.db'; self.cgn = self.root/'cgn.db'
        create_snapshot(self.ae)
        self.file=self.root/'广核.xlsx'
        book=Workbook(); ws=book.active
        self.headers=['质量计划名称','质量计划编码','版本','状态','创建日期','项目','机组','供应商','内部编码','状态','新增加的业务信息']
        ws.append(self.headers)
        ws.append(['广核计划','HYT-000','A','已审查','2026-10-23','HYT','HYT5/6','测试供方甲','CGN-001','CFC','完整保留'])
        ws.append(['广核计划','HYT-000','B','审查中','2026-10-23','HYT','HYT5/6','测试供方甲','CGN-002','历史数据','另一版本'])
        book.save(self.file); book.close()
        import_excel(self.cgn,self.file)
        self.env=patch.dict(os.environ,{'ECS_DB_PATH':str(self.ae),'CGN_DB_PATH':str(self.cgn),'PLAN_CATALOG_PATH':str(self.root/'catalog.db'),'LAN_SYNC_TOKEN':'test-sync','DASHBOARD_ADMIN_TOKEN':'test-admin'})
        self.env.start(); self.client=TestClient(app)

    def tearDown(self):
        self.client.close(); self.env.stop(); self.tmp.cleanup()

    def get(self,path,**params):
        response=self.client.get(path,params=params); self.assertEqual(200,response.status_code,response.text); return response.json()

    def test_combined_filter_and_duplicate_codes_keep_versions(self):
        self.assertEqual(44,self.get('/api/documents')['total'])
        cgn=self.get('/api/documents',source='cgn')
        self.assertEqual(2,cgn['total']); self.assertEqual({'HYT-000'},{r['displayCode'] for r in cgn['items']})
        self.assertEqual(2,len({r['code'] for r in cgn['items']}))
        self.assertEqual({'cgn'},{r['source'] for r in cgn['items']})
        self.assertEqual(42,self.get('/api/documents',source='ae')['total'])
        self.assertEqual(2,self.get('/api/documents',source='cgn',q='HYT-000')['total'])
        self.assertEqual(400,self.client.get('/api/documents?source=other').status_code)

    def test_reviewed_does_not_mean_executed_and_source_fields_retained(self):
        overview=self.get('/api/overview',source='cgn')
        self.assertEqual('0',overview['metrics'][1]['value'])
        self.assertEqual(1,self.get('/api/documents',source='cgn',status='处理中')['total'])
        self.assertEqual(1,self.get('/api/documents',source='cgn',status='完成情况待确认')['total'])
        item=self.get('/api/documents',source='cgn')['items'][0]
        detail=self.get('/api/documents/detail',code=item['code'],source='cgn')
        self.assertEqual(11,len(detail['fields']))
        self.assertEqual(['流程状态','版本状态'],[f['label'] for f in detail['fields'] if f['name']=='状态'])
        self.assertIn('新增加的业务信息',[f['name'] for f in detail['fields']])
        self.assertEqual(404,self.client.get('/api/documents/detail',params={'source':'ae','code':item['code']}).status_code)
        quality=self.get('/api/governance',source='cgn')
        report=next(f for f in quality['fillRates'] if f['field']=='核安全报送级别')
        self.assertIsNone(report['rate']); self.assertEqual(2,report['notProvided'])

    def test_changes_and_notes_do_not_cross_platforms(self):
        book=load_workbook(self.file); book.active.cell(3,4,'已审查'); book.save(self.file); book.close()
        stats=import_excel(self.cgn,self.file); self.assertEqual(1,stats['changed_fields'])
        item=next(r for r in self.get('/api/documents',source='cgn')['items'] if r['version']=='B')
        self.client.put('/api/documents/note',headers={'Authorization':'Bearer test-admin'},json={'code':item['code'],'internalCode':'MANUAL'})
        self.assertNotEqual('MANUAL',self.get('/api/documents',source='ae',q='HYT-000')['items'][0]['internalCode'])
        quality=self.get('/api/governance',source='cgn')
        self.assertEqual(1,len(quality['recentChanges'])); self.assertEqual('cgn',quality['recentChanges'][0]['source'])
        with closing(sqlite3.connect(self.cgn)) as conn:
            self.assertEqual(2,conn.execute('SELECT COUNT(*) FROM documents').fetchone()[0])

    def test_source_push_and_legacy_push_are_independent(self):
        auth={'Authorization':'Bearer test-sync'}
        ae_before=self.ae.read_bytes(); cgn_before=self.cgn.read_bytes()
        self.assertEqual(200,self.client.put('/api/sync/database?source=cgn',headers=auth,content=cgn_before).status_code)
        self.assertEqual(ae_before,self.ae.read_bytes())
        self.assertEqual(400,self.client.put('/api/sync/database?source=cgn',headers=auth,content=ae_before).status_code)
        self.assertEqual(400,self.client.put('/api/sync/database',headers=auth,content=cgn_before).status_code)
        self.assertEqual(200,self.client.put('/api/sync/database',headers=auth,content=ae_before).status_code)
        self.assertEqual(cgn_before,self.cgn.read_bytes())
        status=self.get('/api/sync/status'); self.assertEqual(44,status['documentCount'])
        self.assertEqual(2,len(status['sources']))

    def test_catalog_is_platform_scoped_and_multi_version_requires_review(self):
        mapping={k:i for i,k in enumerate(catalog.FIELDS)}
        catalog.import_rows([['AE计划','HYT-000','','HYT','HYT5/6','','','','ae'], ['广核计划','HYT-000','','HYT','HYT5/6','','','','cgn']],mapping,'计划.csv')
        self.assertEqual({'已完成':1},self.get('/api/catalog',source='ae')['summary'])
        self.assertEqual({'待核对':1},self.get('/api/catalog',source='cgn')['summary'])
        self.assertEqual(50,self.get('/api/catalog')['completionRate'])

    def test_incorrect_excel_dimensions_do_not_truncate_rows(self):
        original=self.file.read_bytes()
        from io import BytesIO
        with ZipFile(BytesIO(original)) as src, ZipFile(self.root/'wrong-dimension.xlsx','w',ZIP_DEFLATED) as dst:
            for name in src.namelist():
                content=src.read(name)
                if name=='xl/worksheets/sheet1.xml':
                    content=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:K2"',content)
                dst.writestr(name,content)
        result=import_excel(self.root/'dimensions.db',self.root/'wrong-dimension.xlsx',expected_count=2)
        self.assertEqual(2,result['inserted'])

    def test_scheduler_continues_cgn_when_ae_fails(self):
        import scheduler
        config={'sources':{'ae':True,'cgn':True},'intervalMinutes':60}
        from argparse import Namespace
        with patch.object(scheduler,'parse_args',return_value=Namespace(once=True,phone=None,force_login=False)), patch.object(scheduler,'load_scheduler_config',return_value=config), patch.object(scheduler,'load_atrust_login_args',side_effect=RuntimeError('AE 配置不可用')), patch('cgn_sync.sync_once') as cgn:
            self.assertEqual(1,scheduler.main())
            cgn.assert_called_once()

    def test_import_count_mismatch_is_atomic(self):
        before=self.cgn.read_bytes()
        with self.assertRaises(ValueError): import_excel(self.cgn,self.file,expected_count=999)
        self.assertEqual(before,self.cgn.read_bytes())

if __name__=='__main__': unittest.main()
